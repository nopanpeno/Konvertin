"""Tes kompres DOCX/PPTX (PRD C5): gambar di dalam dokumen diperkecil, sisanya tidak berubah."""
import io
import zipfile

from PIL import Image

from test_api import _post, _wait


def _media_sizes(data: bytes, prefix: str) -> dict[str, int]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {i.filename: i.file_size for i in z.infolist() if i.filename.startswith(prefix)}


def _run(client, data: bytes, name: str, jtype: str, level: str | None = None):
    opts = {"options": f'{{"level": "{level}"}}'} if level else {}
    r = client.post("/api/jobs", files={"file": (name, data)}, data={"type": jtype, **opts})
    assert r.status_code == 202, r.text
    s = _wait(client, r.json()["job_id"])
    assert s["status"] == "done", s
    return s, client.get(s["download_url"]).content


def test_compress_docx_shrinks_image_and_keeps_document(client, docx_with_image):
    s, out = _run(client, docx_with_image, "Laporan Foto.docx", "compress_docx", "max")
    assert s["download_name"] == "Laporan Foto-kompres.docx"
    assert len(out) < len(docx_with_image) * 0.5

    from docx import Document

    doc = Document(io.BytesIO(out))  # masih dokumen Word yang valid
    assert "Dokumen dengan foto besar." in [p.text for p in doc.paragraphs]
    assert len(doc.inline_shapes) == 1  # gambar tetap ada di tempatnya
    part = next(p for p in doc.part.package.iter_parts() if p.partname.startswith("/word/media/"))
    assert max(Image.open(io.BytesIO(part.blob)).size) <= 1024  # batas level maksimal


def test_compress_levels_order(client, docx_with_image):
    sizes = {lv: len(_run(client, docx_with_image, "a.docx", "compress_docx", lv)[1]) for lv in ("light", "medium", "max")}
    assert sizes["max"] < sizes["medium"] < sizes["light"] < len(docx_with_image)


def test_compress_pptx_keeps_names_alpha_and_non_images(client, pptx_with_image):
    s, out = _run(client, pptx_with_image, "Slide Saya.pptx", "compress_pptx", "medium")
    assert s["download_name"] == "Slide Saya-kompres.pptx"
    assert len(out) < len(pptx_with_image)

    with zipfile.ZipFile(io.BytesIO(pptx_with_image)) as a, zipfile.ZipFile(io.BytesIO(out)) as b:
        assert a.namelist() == b.namelist()  # nama entri dan urutan sama: relasi tidak putus
        assert b.read("ppt/media/clip1.mp4") == a.read("ppt/media/clip1.mp4")  # non-gambar utuh
        assert b.read("ppt/presentation.xml") == a.read("ppt/presentation.xml")
        png = Image.open(io.BytesIO(b.read("ppt/media/image2.png")))
        assert png.mode == "RGBA" and png.getpixel((10, 10))[3] == 0  # transparansi tetap
        assert png.getpixel((png.width - 10, 10))[3] == 255
        assert max(png.size) <= 1600
        assert Image.open(io.BytesIO(b.read("ppt/media/image1.jpeg"))).format == "JPEG"


def test_compress_docx_without_images_stays_valid_and_not_larger(client, sample_docx):
    s, out = _run(client, sample_docx, "polos.docx", "compress_docx")
    assert len(out) <= len(sample_docx)  # tidak ada gambar: paling banyak hemat sedikit dari XML

    from docx import Document

    text = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
    assert "Pasal 1" in text
    with zipfile.ZipFile(io.BytesIO(sample_docx)) as a, zipfile.ZipFile(io.BytesIO(out)) as b:
        assert a.namelist() == b.namelist()
        assert all(a.read(n) == b.read(n) for n in a.namelist())  # isi tiap bagian identik


def test_compress_never_larger_than_input(client, docx_with_image):
    # kualitas sudah ringan pun, hasil tidak boleh melebihi aslinya
    _, out = _run(client, docx_with_image, "a.docx", "compress_docx", "light")
    assert len(out) <= len(docx_with_image)


def test_compress_rejects_wrong_kind(client, sample_pdf, docx_with_image):
    r = _post(client, sample_pdf, "a.pdf", "compress_docx")
    assert r.status_code == 400 and "PDF" in r.json()["error"]
    r = _post(client, docx_with_image, "a.docx", "compress_pptx")
    assert r.status_code == 400 and "DOCX" in r.json()["error"]


def test_compress_bad_level_falls_back_to_medium():
    from app.validation import check_options

    assert check_options("compress_docx", {"level": "../etc"}) == {"level": "medium"}
    assert check_options("compress_pptx", {"level": "max"}) == {"level": "max"}


def test_detect_offers_compress_for_office_files(client, docx_with_image, pptx_with_image):
    for data, kind, job in ((docx_with_image, "docx", "compress_docx"), (pptx_with_image, "pptx", "compress_pptx")):
        r = client.post("/api/detect", files={"file": ("x.bin", data)}).json()
        assert r["kind"] == kind
        assert {"type": job, "output": kind} in r["options"]


def test_conversions_list_hides_compress_jobs(client):
    c = client.get("/api/conversions").json()["conversions"]
    listed = {x["type"] for group in c.values() for x in group}
    assert not {t for t in listed if t.startswith("compress_")}


def test_corrupt_media_is_skipped_not_fatal(tmp_path, big_jpeg):
    """Gambar rusak di dalam dokumen dibiarkan; gambar lain tetap diperkecil."""
    from app import converters

    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    with zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", "<w:document/>")
        z.writestr("word/media/rusak.jpeg", b"bukan gambar" * 100)
        z.writestr("word/media/foto.jpeg", big_jpeg)
    converters.compress_docx(src, dst, {"level": "max"})
    with zipfile.ZipFile(dst) as z:
        assert z.read("word/media/rusak.jpeg") == b"bukan gambar" * 100
        assert len(z.read("word/media/foto.jpeg")) < len(big_jpeg) / 2
