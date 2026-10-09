import io
import os
import sys
import tempfile
from pathlib import Path

_tmp = tempfile.mkdtemp(prefix="konvertin_test_")
os.environ["KONVERTIN_DATA_DIR"] = _tmp
os.environ["KONVERTIN_FRONTEND_DIR"] = str(Path(_tmp) / "nofrontend")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from fastapi.testclient import TestClient


@pytest.fixture(scope="session")
def client():
    from app.main import app

    with TestClient(app) as c:
        yield c


@pytest.fixture(autouse=True)
def _reset_rate():
    from app import ratelimit

    ratelimit.reset()


@pytest.fixture(scope="session")
def sample_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    b = io.BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    c.setFont("Helvetica", 14)
    c.drawString(72, 780, "Laporan Praktikum Konvertin")
    c.setFont("Helvetica", 11)
    for i in range(10):
        c.drawString(72, 740 - i * 18, f"Baris isi nomor {i + 1}: pengujian konversi dokumen.")
    c.showPage()
    c.save()
    return b.getvalue()


@pytest.fixture(scope="session")
def blank_scan_pdf() -> bytes:
    """PDF tanpa teks (meniru hasil scan)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    b = io.BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    c.rect(100, 100, 300, 300, fill=1)
    c.showPage()
    c.save()
    return b.getvalue()


@pytest.fixture(scope="session")
def sample_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Kontrak Kerja", 1)
    d.add_paragraph("Pasal 1: Para pihak sepakat melakukan pengujian.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Nama"
    t.cell(0, 1).text = "Budi"
    t.cell(1, 0).text = "Peran"
    t.cell(1, 1).text = "Penguji"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


@pytest.fixture(scope="session")
def sample_xlsx() -> bytes:
    import zipfile

    # xlsx minimal valid, dibuat tanpa dependensi tambahan
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", '<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr("xl/worksheets/sheet1.xml", '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Total</t></is></c><c r="B1"><v>42</v></c></row></sheetData></worksheet>')
    return b.getvalue()
