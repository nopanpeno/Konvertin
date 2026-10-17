import io
import time
import zipfile
import xml.etree.ElementTree as ET

import pytest

PDF = "application/pdf"


def _post(client, data: bytes, name: str, jtype: str):
    return client.post("/api/jobs", files={"file": (name, data)}, data={"type": jtype})


def _wait(client, job_id: str, timeout=90):
    end = time.time() + timeout
    while time.time() < end:
        r = client.get(f"/api/jobs/{job_id}").json()
        if r["status"] in ("done", "failed"):
            return r
        time.sleep(0.5)
    raise AssertionError("timeout")


def test_health_and_conversions(client):
    assert client.get("/api/health").json() == {"ok": True}
    c = client.get("/api/conversions").json()
    assert {
        kind: {item["type"] for item in conversions}
        for kind, conversions in c["conversions"].items()
    } == {
        "pdf": {"pdf_to_docx", "pdf_to_txt"},
        "docx": {"docx_to_pdf", "docx_to_odt", "docx_to_md", "docx_to_txt"},
        "pptx": {"pptx_to_pdf", "pptx_to_odp", "pptx_to_txt"},
        "xlsx": {"xlsx_to_pdf", "xlsx_to_ods", "xlsx_to_csv"},
        "csv": {"csv_to_xlsx"},
        "txt": {"txt_to_docx"},
    }
    assert c["max_mb"] == 25 and c["max_pages"] == 200


def test_dotenv_files_are_not_served(client, tmp_path, monkeypatch):
    from app import config

    (tmp_path / ".env").write_text("SECRET=must-not-be-served")
    (tmp_path / ".env.example").write_text("SECRET=must-not-be-served")
    monkeypatch.setattr(config, "FRONTEND_DIR", tmp_path)
    assert client.get("/.env").status_code == 404
    assert client.get("/.env.example").status_code == 404
    assert client.get("/nested/.env.local").status_code == 404


@pytest.mark.parametrize(
    "configured,expected",
    [
        ("http://127.0.0.1:8001", "http://127.0.0.1:8001"),
        (
            "http://127.0.0.1:8001/api/auth/google/callback",
            "http://127.0.0.1:8001",
        ),
        (
            "http://127.0.0.1:8001/api/auth/google/callback/",
            "http://127.0.0.1:8001",
        ),
    ],
)
def test_public_base_url_accepts_origin_or_google_callback(configured, expected):
    from app.config import normalize_public_base_url

    assert normalize_public_base_url(configured) == expected


def test_windows_worker_uses_supported_process_options(monkeypatch):
    from app import jobs

    monkeypatch.setattr(jobs.subprocess, "CREATE_NEW_PROCESS_GROUP", 512, raising=False)
    assert jobs._process_options("docx_to_pdf", windows=True) == {"creationflags": 512}


def test_office_conversion_reports_missing_libreoffice(monkeypatch, tmp_path):
    from app import converters
    from app.errors import ConvertError

    monkeypatch.setattr(converters.shutil, "which", lambda _: None)
    monkeypatch.setenv("KONVERTIN_SOFFICE_PATH", "")
    monkeypatch.setenv("PROGRAMFILES", str(tmp_path))
    monkeypatch.setenv("PROGRAMFILES(X86)", str(tmp_path))
    with pytest.raises(ConvertError, match="LibreOffice belum terpasang"):
        converters._soffice_to_pdf(tmp_path / "in.docx", tmp_path / "out.pdf")


def test_detect(client, sample_pdf, sample_docx):
    r = client.post("/api/detect", files={"file": ("x.bin", sample_pdf)}).json()
    # PDF juga menawarkan opsi kompresi (tetap berformat pdf)
    assert r["kind"] == "pdf" and {o["output"] for o in r["options"]} == {"docx", "txt", "pdf"}
    r = client.post("/api/detect", files={"file": ("x.pdf", sample_docx)}).json()
    assert r["kind"] == "docx"


def test_docx_to_pdf(client, sample_docx):
    r = _post(client, sample_docx, "Kontrak Saya.docx", "docx_to_pdf")
    assert r.status_code == 202
    s = _wait(client, r.json()["job_id"])
    assert s["status"] == "done", s
    assert s["download_name"] == "Kontrak Saya.pdf"
    d = client.get(s["download_url"])
    assert d.content.startswith(b"%PDF")


def test_xlsx_to_pdf(client, sample_xlsx):
    s = _wait(client, _post(client, sample_xlsx, "a.xlsx", "xlsx_to_pdf").json()["job_id"])
    assert s["status"] == "done", s


def test_xlsx_to_csv_exports_first_sheet_values(client, sample_xlsx):
    response = _post(client, sample_xlsx, "angka.xlsx", "xlsx_to_csv")
    assert response.status_code == 202
    status = _wait(client, response.json()["job_id"])
    assert status["status"] == "done", status
    result = client.get(status["download_url"]).content.decode("utf-8-sig")
    assert "Total,42" in result


def test_csv_to_xlsx_preserves_text_and_does_not_create_formulas(client):
    response = _post(client, b"Nama,Isi\nAni,=1+1\n", "data.csv", "csv_to_xlsx")
    assert response.status_code == 202
    status = _wait(client, response.json()["job_id"])
    assert status["status"] == "done", status

    with zipfile.ZipFile(io.BytesIO(client.get(status["download_url"]).content)) as package:
        sheet = package.read("xl/worksheets/sheet1.xml")
    root = ET.fromstring(sheet)
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cells = root.findall(".//x:c", ns)
    assert cells[3].get("t") == "inlineStr"
    assert cells[3].find("x:f", ns) is None
    assert cells[3].find("x:is/x:t", ns).text == "=1+1"
    converted = client.get(status["download_url"]).content
    pdf_response = _post(client, converted, "data.xlsx", "xlsx_to_pdf")
    pdf_status = _wait(client, pdf_response.json()["job_id"])
    assert pdf_status["status"] == "done", pdf_status


def test_txt_to_docx_and_docx_to_markdown_are_supported(client, sample_docx):
    response = _post(client, sample_docx, "catatan.docx", "docx_to_md")
    status = _wait(client, response.json()["job_id"])
    assert status["status"] == "done", status
    markdown = client.get(status["download_url"]).text
    assert "# Kontrak Kerja" in markdown and "Pasal 1" in markdown

    response = _post(client, "Baris pertama\nBaris kedua\n".encode(), "catatan.txt", "txt_to_docx")
    assert response.status_code == 202
    status = _wait(client, response.json()["job_id"])
    assert status["status"] == "done", status
    from docx import Document

    document = Document(io.BytesIO(client.get(status["download_url"]).content))
    assert [paragraph.text for paragraph in document.paragraphs] == ["Baris pertama", "Baris kedua"]


def test_pptx_to_txt_exports_text_by_slide(client):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "Ringkasan presentasi"
    source = io.BytesIO()
    presentation.save(source)

    response = _post(client, source.getvalue(), "presentasi.pptx", "pptx_to_txt")
    assert response.status_code == 202
    status = _wait(client, response.json()["job_id"])
    assert status["status"] == "done", status
    assert "Slide 1" in client.get(status["download_url"]).text
    assert "Ringkasan presentasi" in client.get(status["download_url"]).text


def test_office_files_export_to_open_document_formats(client, sample_docx, sample_xlsx):
    docx_status = _wait(client, _post(client, sample_docx, "dokumen.docx", "docx_to_odt").json()["job_id"])
    xlsx_status = _wait(client, _post(client, sample_xlsx, "lembar.xlsx", "xlsx_to_ods").json()["job_id"])
    assert docx_status["status"] == "done", docx_status
    assert xlsx_status["status"] == "done", xlsx_status
    assert zipfile.is_zipfile(io.BytesIO(client.get(docx_status["download_url"]).content))
    assert zipfile.is_zipfile(io.BytesIO(client.get(xlsx_status["download_url"]).content))


def test_pptx_to_pdf(client):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    text_box.text = "Konvertin PowerPoint"
    source = io.BytesIO()
    presentation.save(source)

    response = _post(client, source.getvalue(), "presentasi.pptx", "pptx_to_pdf")
    assert response.status_code == 202
    status = _wait(client, response.json()["job_id"])
    assert status["status"] == "done", status
    assert status["download_name"] == "presentasi.pdf"
    from pypdf import PdfReader

    result = client.get(status["download_url"])
    assert result.content.startswith(b"%PDF")
    pdf = PdfReader(io.BytesIO(result.content))
    assert len(pdf.pages) == 1
    assert "Konvertin PowerPoint" in pdf.pages[0].extract_text()


def test_pptx_to_odp_exports_open_document_presentation(client):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Slide ODP"
    source = io.BytesIO()
    presentation.save(source)
    response = _post(client, source.getvalue(), "presentasi.pptx", "pptx_to_odp")
    assert response.status_code == 202
    status = _wait(client, response.json()["job_id"])
    assert status["status"] == "done", status
    assert zipfile.is_zipfile(io.BytesIO(client.get(status["download_url"]).content))


def test_pdf_to_docx(client, sample_pdf):
    s = _wait(client, _post(client, sample_pdf, "laporan.pdf", "pdf_to_docx").json()["job_id"])
    assert s["status"] == "done", s
    from docx import Document

    doc = Document(io.BytesIO(client.get(s["download_url"]).content))
    assert "Laporan Praktikum Konvertin" in "\n".join(p.text for p in doc.paragraphs)


def test_compress_pdf_shrinks_embedded_image_and_keeps_text(client, big_jpeg):
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    source = io.BytesIO()
    pdf = canvas.Canvas(source)
    pdf.drawString(72, 780, "Laporan tetap terbaca")
    pdf.drawImage(ImageReader(io.BytesIO(big_jpeg)), 36, 300, width=520, height=390)
    pdf.save()
    original = source.getvalue()

    response = _post(client, original, "laporan.pdf", "compress_pdf")
    assert response.status_code == 202
    status = _wait(client, response.json()["job_id"])
    assert status["status"] == "done", status
    assert status["download_name"] == "laporan-kompres.pdf"
    result = client.get(status["download_url"]).content

    from pypdf import PdfReader

    assert len(result) < len(original)
    assert "Laporan tetap terbaca" in PdfReader(io.BytesIO(result)).pages[0].extract_text()


def test_pdf_to_txt(client, sample_pdf):
    s = _wait(client, _post(client, sample_pdf, "laporan.pdf", "pdf_to_txt").json()["job_id"])
    assert s["status"] == "done", s
    assert "Baris isi nomor 3" in client.get(s["download_url"]).text


def test_docx_to_txt(client, sample_docx):
    s = _wait(client, _post(client, sample_docx, "k.docx", "docx_to_txt").json()["job_id"])
    assert s["status"] == "done", s
    t = client.get(s["download_url"]).text
    assert "Pasal 1" in t and "Budi" in t


def test_scan_pdf_rejected_with_honest_message(client, blank_scan_pdf):
    for jt in ("pdf_to_docx", "pdf_to_txt"):
        s = _wait(client, _post(client, blank_scan_pdf, "scan.pdf", jt).json()["job_id"])
        assert s["status"] == "failed" and "scan" in s["error"].lower() and "OCR" in s["error"]


def test_magic_bytes_not_extension(client, sample_docx):
    # DOCX diberi nama .pdf dan dipakai untuk pdf_to_docx -> ditolak
    r = _post(client, sample_docx, "palsu.pdf", "pdf_to_docx")
    assert r.status_code == 400 and "DOCX" in r.json()["error"]


def test_garbage_rejected(client):
    r = _post(client, b"hello world, bukan dokumen", "a.pdf", "pdf_to_docx")
    assert r.status_code == 400


def test_empty_rejected(client):
    assert _post(client, b"", "a.pdf", "pdf_to_docx").status_code == 400


def test_unknown_type(client, sample_pdf):
    assert _post(client, sample_pdf, "a.pdf", "pdf_to_mp3").status_code == 400


def test_encrypted_pdf_rejected(client, sample_pdf):
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter()
    for p in PdfReader(io.BytesIO(sample_pdf)).pages:
        w.add_page(p)
    w.encrypt("rahasia")
    b = io.BytesIO()
    w.write(b)
    r = _post(client, b.getvalue(), "k.pdf", "pdf_to_docx")
    assert r.status_code == 400 and "password" in r.json()["error"]


def test_old_ole_format_rejected(client):
    r = _post(client, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600, "a.docx", "docx_to_pdf")
    assert r.status_code == 400 and "password" in r.json()["error"]


def test_zip_bomb_rejected(client):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("word/document.xml", b"\0" * (60 * 1024 * 1024))
    r = _post(client, b.getvalue(), "bom.docx", "docx_to_pdf")
    assert r.status_code == 400


def test_too_large(client, monkeypatch):
    from app import config

    monkeypatch.setattr(config, "MAX_DOC_BYTES", 1024)
    r = _post(client, b"%PDF-" + b"0" * 5000, "a.pdf", "pdf_to_txt")
    assert r.status_code == 413


def test_rate_limit(client, sample_pdf, monkeypatch):
    from app import config

    monkeypatch.setattr(config, "RATE_LIMIT_JOBS", 2)
    codes = [_post(client, b"junk", "a.pdf", "pdf_to_txt").status_code for _ in range(3)]
    assert codes == [400, 400, 429]


def test_cancel_and_not_found(client, sample_pdf):
    jid = _post(client, sample_pdf, "a.pdf", "pdf_to_txt").json()["job_id"]
    assert client.delete(f"/api/jobs/{jid}").status_code == 200
    assert client.get(f"/api/jobs/{jid}/download").status_code == 404
    assert client.get("/api/jobs/tidak-ada").status_code == 404


def test_cleanup_deletes_files(client, sample_docx):
    from app import config, jobs

    s = _wait(client, _post(client, sample_docx, "k.docx", "docx_to_txt").json()["job_id"])
    assert s["status"] == "done"
    assert any(config.STORAGE_DIR.iterdir())
    jobs._set(s["job_id"], expires_at=time.time() - 1)
    jobs.cleanup_expired()
    assert client.get(f"/api/jobs/{s['job_id']}").status_code == 404
    assert not list(config.STORAGE_DIR.glob("*.out")) or all(
        f.name != jobs.get(s["job_id"]) for f in config.STORAGE_DIR.iterdir()
    )


def test_email_registration_requires_verification_and_logout_invalidates_session(client, monkeypatch):
    from urllib.parse import parse_qs, urlsplit

    from app import accounts, config, mail

    client.cookies.clear()
    monkeypatch.setattr(config, "PUBLIC_BASE_URL", "https://konvertin.test")
    monkeypatch.setattr(config, "SESSION_COOKIE_SECURE", False)
    monkeypatch.setattr(config, "SMTP_HOST", "smtp.example.test")
    monkeypatch.setattr(config, "SMTP_FROM_EMAIL", "noreply@konvertin.test")
    sent = []
    monkeypatch.setattr(mail, "send_verification", lambda email, name, url: sent.append((email, name, url)))
    response = client.post("/api/auth/register", json={
        "email": "verify-flow@example.test", "password": "correct-horse-2026",
        "display_name": "Verify Flow", "phone": "+62 812-3456-7890",
    })
    assert response.status_code == 202
    assert "set-cookie" not in response.headers
    assert sent and sent[0][0] == "verify-flow@example.test"
    assert client.post("/api/auth/login", json={
        "email": "verify-flow@example.test", "password": "correct-horse-2026",
    }).status_code == 403
    token = parse_qs(urlsplit(sent[0][2]).query)["token"][0]
    verified = client.get("/api/auth/verify-email", params={"token": token}, follow_redirects=False)
    assert verified.status_code == 303
    assert verified.headers["location"] == "https://konvertin.test/akun.html?verified=1"
    login = client.post("/api/auth/login", json={
        "email": "verify-flow@example.test", "password": "correct-horse-2026",
    })
    assert login.status_code == 200
    assert login.json()["user"]["email_verified"] is True
    assert login.json()["user"]["phone"] == "+62 812-3456-7890"
    logout = client.post("/api/auth/logout")
    assert logout.status_code == 200
    assert "max-age=0" in logout.headers["set-cookie"].lower()
    assert client.get("/api/auth/me").json()["user"] is None
    assert not accounts.verify_email_token(token)


def test_admin_user_detail_warning_controls_and_account_erasure(client, monkeypatch):
    from app import accounts

    client.cookies.clear()
    user_id, error = accounts.create_user(
        "managed-user@example.test", "correct-horse-2026", "Managed User", "+62 812-1234-5678",
    )
    assert error is None
    with accounts._db() as c:
        c.execute("UPDATE users SET email_verified=1 WHERE user_id=?", (user_id,))
        c.execute("""INSERT INTO payments(user_id,external_id,amount,currency,status,created_at)
            VALUES(?,?,?,'IDR','paid',?)""", (user_id, "managed-payment", 25000, time.time()))
    accounts.record_browser_usage(user_id, "jpg-ke-png")

    admin_id, error = accounts.create_user(
        "dashboard-admin@example.test", "another-correct-horse", "Dashboard Admin",
    )
    assert error is None
    with accounts._db() as c:
        c.execute("UPDATE users SET role='superadmin',email_verified=1 WHERE user_id=?", (admin_id,))
    token = accounts.new_session(admin_id)
    client.cookies.set(accounts.COOKIE, token)

    details = client.get(f"/api/admin/users/{user_id}")
    assert details.status_code == 200
    assert details.json()["user"]["phone"] == "+62 812-1234-5678"
    assert details.json()["history"][0]["tool_type"] == "jpg-ke-png"

    warning = client.post(f"/api/admin/users/{user_id}/notifications", json={
        "kind": "critical", "title": "Peringatan akun", "message": "Harap tinjau aktivitas akun.",
    })
    assert warning.status_code == 201
    assert warning.json()["email_status"] == "not_configured"
    prefs = client.patch(f"/api/admin/users/{user_id}", json={"email_notifications": False})
    assert prefs.status_code == 200
    assert client.get("/api/admin/users?search=managed-user").json()["users"][0]["email_notifications"] is False

    deleted = client.delete(f"/api/admin/users/{user_id}")
    assert deleted.status_code == 200 and deleted.json()["payments_anonymized"] is True
    with accounts._db() as c:
        assert c.execute("SELECT 1 FROM users WHERE user_id=?", (user_id,)).fetchone() is None
        payment = c.execute("SELECT user_id,external_id,invoice_url FROM payments WHERE amount=25000").fetchone()
        assert payment["user_id"] == 0 and payment["external_id"].startswith("deleted-")
        assert payment["invoice_url"] is None
        usage = c.execute("SELECT user_id FROM usage_events WHERE tool_type='jpg-ke-png'").fetchone()
        assert usage["user_id"] is None

    assert client.delete(f"/api/admin/users/{admin_id}").status_code == 403


def test_staff_admin_cannot_change_sensitive_settings(client):
    from app import accounts

    client.cookies.clear()
    staff_id, error = accounts.create_user(
        "staff-admin@example.test", "staff-secret-password", "Staff Admin",
    )
    assert error is None
    with accounts._db() as c:
        c.execute("UPDATE users SET role='admin',email_verified=1 WHERE user_id=?", (staff_id,))
    client.cookies.set(accounts.COOKIE, accounts.new_session(staff_id))

    assert client.get("/api/admin/dashboard").status_code == 200
    assert client.get("/api/admin/dashboard").json().get("revenue_last_30_days") is None
    assert client.patch("/api/admin/settings", json={"free_daily_quota": 30}).status_code == 403
    assert client.get("/api/admin/audit").status_code == 403
    assert client.get(f"/api/admin/users/{staff_id}").status_code == 403


def test_google_login_creates_users_and_promotes_allowlisted_owners(client, monkeypatch):
    from app import accounts, config, google_auth
    from urllib.parse import parse_qs, urlsplit

    client.cookies.clear()
    monkeypatch.setattr(config, "GOOGLE_CLIENT_ID", "test-google-client")
    monkeypatch.setattr(config, "GOOGLE_CLIENT_SECRET", "test-google-secret")
    monkeypatch.setattr(config, "PUBLIC_BASE_URL", "https://konvertin.example")
    monkeypatch.setattr(config, "SESSION_COOKIE_SECURE", False)
    started = client.get("/api/auth/google", follow_redirects=False)
    assert started.status_code == 302
    assert "accounts.google.com" in started.headers["location"]
    assert "code_challenge_method=S256" in started.headers["location"]
    assert "konvertin_google_nonce" in started.headers["set-cookie"]

    monkeypatch.setattr(google_auth, "verify_authorization_code", lambda code, verifier, nonce: {
        "email": "user@gmail.com", "display_name": "Pengguna Google",
    })
    wrong_state = client.get("/api/auth/google/callback?code=code&state=invalid", follow_redirects=False)
    assert wrong_state.status_code == 400
    state = parse_qs(urlsplit(started.headers["location"]).query)["state"][0]
    callback = client.get(f"/api/auth/google/callback?code=code&state={state}", follow_redirects=False)
    assert callback.status_code == 303, callback.text
    assert callback.headers["location"] == "https://konvertin.example/profil.html"
    owner = client.get("/api/auth/me").json()["user"]
    assert owner["email"] == "user@gmail.com"
    assert owner["role"] == "user"

    client.cookies.clear()
    started = client.get("/api/auth/google", follow_redirects=False)
    state = parse_qs(urlsplit(started.headers["location"]).query)["state"][0]
    monkeypatch.setattr(google_auth, "verify_authorization_code", lambda code, verifier, nonce: {
        "email": "setara.sea@gmail.com", "display_name": "Pemilik Konvertin",
    })
    callback = client.get(f"/api/auth/google/callback?code=code&state={state}", follow_redirects=False)
    assert callback.status_code == 303, callback.text
    assert callback.headers["location"] == "https://konvertin.example/admin.html"
    owner = client.get("/api/auth/me").json()["user"]
    assert owner["role"] == "superadmin"
    assert client.patch(f"/api/admin/users/{owner['id']}", json={"status": "suspended"}).status_code == 403

    assert accounts.authenticate("setara.sea@gmail.com", "any-password") is None
    assert accounts.google_account("unauthorized@example.test", "Unauthorized")["role"] == "user"


def test_admin_login_explains_missing_secret_and_bootstraps_when_configured(client, monkeypatch):
    from app import accounts, config

    client.cookies.clear()
    monkeypatch.setattr(config, "ADMIN_EMAIL", "admin@konvertin.com")
    monkeypatch.setattr(config, "ADMIN_PASSWORD", "")
    response = client.post("/api/auth/login", json={
        "email": config.ADMIN_EMAIL, "password": "not-the-secret",
    })
    assert response.status_code == 503
    assert "KONVERTIN_ADMIN_PASSWORD" in response.json()["error"]

    secret = "unit-test-admin-secret-2026"
    monkeypatch.setattr(config, "ADMIN_PASSWORD", secret)
    accounts.init()
    try:
        google_admin = accounts.google_account(config.ADMIN_EMAIL, "Admin Google")
        assert google_admin["role"] == "admin"
        response = client.post("/api/auth/login", json={
            "email": config.ADMIN_EMAIL, "password": secret,
        })
        assert response.status_code == 200, response.text
        assert response.json()["user"]["role"] == "admin"
    finally:
        client.cookies.clear()
        with accounts._db() as c:
            admin = c.execute("SELECT user_id FROM users WHERE email=?", (config.ADMIN_EMAIL,)).fetchone()
            if admin:
                c.execute("DELETE FROM sessions WHERE user_id=?", (admin["user_id"],))
            c.execute("DELETE FROM users WHERE email=?", (config.ADMIN_EMAIL,))


def test_google_token_validation_checks_audience_email_and_nonce(monkeypatch):
    import base64
    import json

    from app import config, google_auth

    monkeypatch.setattr(config, "GOOGLE_CLIENT_ID", "expected-client")
    monkeypatch.setattr(config, "GOOGLE_CLIENT_SECRET", "oauth-secret")
    monkeypatch.setattr(config, "PUBLIC_BASE_URL", "https://konvertin.example")
    payload = base64.urlsafe_b64encode(json.dumps({"nonce": "expected-nonce"}).encode()).rstrip(b"=").decode()
    id_token = f"e30.{payload}.signature"
    responses = iter([
        {"id_token": id_token},
        {
            "aud": "expected-client", "iss": "https://accounts.google.com",
            "exp": int(time.time()) + 300, "email_verified": "true",
            "email": "kartikafajar13@gmail.com", "name": "Pemilik Kedua",
        },
    ])
    monkeypatch.setattr(google_auth, "_json_response", lambda request: next(responses))
    identity = google_auth.verify_authorization_code("oauth-code", "pkce-verifier", "expected-nonce")
    assert identity == {"email": "kartikafajar13@gmail.com", "display_name": "Pemilik Kedua"}

    responses = iter([
        {"id_token": id_token},
        {
            "aud": "unexpected-client", "iss": "https://accounts.google.com",
            "exp": int(time.time()) + 300, "email_verified": "true",
            "email": "kartikafajar13@gmail.com",
        },
    ])
    monkeypatch.setattr(google_auth, "_json_response", lambda request: next(responses))
    with pytest.raises(ValueError, match="tidak terverifikasi"):
        google_auth.verify_authorization_code("oauth-code", "pkce-verifier", "expected-nonce")

    responses = iter([
        {"id_token": id_token},
        {
            "aud": "expected-client", "iss": "accounts.google.com",
            "exp": int(time.time()) + 300, "email_verified": "true",
            "email": "ordinary-user@example.com", "name": "Pengguna Google",
        },
    ])
    monkeypatch.setattr(google_auth, "_json_response", lambda request: next(responses))
    identity = google_auth.verify_authorization_code("oauth-code", "pkce-verifier", "expected-nonce")
    assert identity == {"email": "ordinary-user@example.com", "display_name": "Pengguna Google"}


def test_midtrans_webhook_activates_subscription_once(client, monkeypatch):
    import hashlib

    from app import accounts, config

    client.cookies.clear()
    monkeypatch.setattr(config, "MIDTRANS_SERVER_KEY", "midtrans-test-key")
    user_id, error = accounts.create_user("midtrans-user@example.test", "user-password-2026", "Midtrans User")
    assert error is None
    now = time.time()
    with accounts._db() as c:
        c.execute("UPDATE users SET email_verified=1 WHERE user_id=?", (user_id,))
        c.execute("""INSERT INTO subscriptions(user_id,plan_id,status,created_at)
            VALUES(?,'premium_monthly','pending',?)""", (user_id, now))
        subscription_id = c.execute("SELECT last_insert_rowid()").fetchone()[0]
        c.execute("""INSERT INTO payments(user_id,subscription_id,external_id,amount,currency,status,provider,created_at)
            VALUES(?,?,? ,1000,'IDR','pending','midtrans',?)""",
                  (user_id, subscription_id, "KVT-test-order", now))
    gross_amount = "1000.00"
    signature = hashlib.sha512(
        f"KVT-test-order200{gross_amount}{config.MIDTRANS_SERVER_KEY}".encode()
    ).hexdigest()
    payload = {
        "order_id": "KVT-test-order", "status_code": "200", "gross_amount": gross_amount,
        "signature_key": signature, "transaction_status": "settlement", "transaction_id": "trx-1",
    }
    assert client.post("/api/billing/midtrans/webhook", json=payload).status_code == 200
    end_time = accounts.account_summary(user_id)["subscription"]["ends_at"]
    assert end_time > now
    assert client.post("/api/billing/midtrans/webhook", json=payload).status_code == 200
    with accounts._db() as c:
        assert c.execute("SELECT status FROM payments WHERE external_id='KVT-test-order'").fetchone()["status"] == "paid"
        assert c.execute("SELECT count(*) FROM user_notifications WHERE user_id=? AND title='Pembayaran berhasil'",
                         (user_id,)).fetchone()[0] == 1
        assert c.execute("SELECT ends_at FROM subscriptions WHERE subscription_id=?",
                         (subscription_id,)).fetchone()["ends_at"] == end_time


def test_expired_subscription_reverts_to_free_and_notifies_once(client):
    from app import accounts

    user_id, error = accounts.create_user("expired-user@example.test", "user-password-2026", "Expired User")
    assert error is None
    with accounts._db() as c:
        c.execute("UPDATE users SET email_verified=1 WHERE user_id=?", (user_id,))
        c.execute("""INSERT INTO subscriptions(user_id,plan_id,status,started_at,ends_at,created_at)
            VALUES(?,'premium_monthly','active',?,?,?)""", (user_id, time.time() - 3600, time.time() - 1, time.time() - 3600))

    summary = accounts.account_summary(user_id)
    assert summary["subscription"] is None
    assert summary["plan"]["plan_id"] == "free"
    accounts.account_summary(user_id)
    with accounts._db() as c:
        assert c.execute("SELECT status FROM subscriptions WHERE user_id=?", (user_id,)).fetchone()["status"] == "expired"
        assert c.execute("SELECT count(*) FROM user_notifications WHERE user_id=? AND title='Langganan Premium berakhir'",
                         (user_id,)).fetchone()[0] == 1
