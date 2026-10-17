"""Konverter dokumen. Dijalankan sebagai proses terpisah (lihat jobs.py) supaya
bisa dibatalkan, di-timeout, dan dibatasi memorinya.

CLI:  python -m app.converters <job_type> <input> <output> [options_json]
Exit code 0 = sukses. Pada gagal, pesan untuk user ditulis ke stderr dengan awalan 'USER:'.
"""
import csv
import inspect
import io
import json
import os
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from .errors import ConvertError

SCAN_MESSAGE = (
    "PDF ini tampaknya hasil scan (isinya gambar, tidak ada teks yang bisa dibaca). "
    "PDF hasil scan butuh OCR dan belum didukung di versi ini."
)


def _has_text(pdf: Path, max_pages: int = 10, min_chars: int = 3) -> bool:
    from pypdf import PdfReader

    reader = PdfReader(str(pdf))
    total = 0
    for page in list(reader.pages)[:max_pages]:
        total += len((page.extract_text() or "").strip())
        if total >= min_chars:
            return True
    return False


# level -> (kualitas gambar 1-100 untuk downsampling, dpi target)
COMPRESS_LEVELS = {"light": (80, 150), "medium": (60, 120), "max": (40, 90)}


def _resize_pixmap(pix, mat):
    """Perkecil Pixmap dengan rasio terjaga lewat render ke halaman sementara."""
    import pymupdf as fitz

    tmp = fitz.open()
    page = tmp.new_page(width=pix.width, height=pix.height)
    page.insert_image(page.rect, pixmap=pix)
    out = page.get_pixmap(matrix=mat, alpha=False)
    tmp.close()
    return out


def compress_pdf(src: Path, dst: Path, options: dict | None = None) -> None:
    """Kecilkan PDF dengan menurunkan resolusi gambar di dalamnya.

    Tiga level (PRD C2). Kalau hasil malah lebih besar dari aslinya,
    file asli yang disalin ke output supaya user tidak rugi.
    """
    import pymupdf as fitz  # PyMuPDF: buka, ubah gambar, simpan

    level = (options or {}).get("level", "medium")
    quality, dpi = COMPRESS_LEVELS.get(level, COMPRESS_LEVELS["medium"])
    doc = fitz.open(str(src))
    try:
        seen: set[int] = set()
        for page in doc:
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen:  # gambar yang dipakai banyak halaman cukup sekali
                    continue
                seen.add(xref)
                if img[1]:  # punya mask transparan: jangan diganggu
                    continue
                base = doc.extract_image(xref)
                if base["ext"].lower() not in ("jpeg", "jpg", "png"):
                    continue
                w, h = base["width"], base["height"]
                target_w = max(1, int(round(w * dpi / 300)))
                pix = fitz.Pixmap(doc, xref)
                if pix.colorspace is None or pix.colorspace.n > 3 or pix.alpha:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                jpg = pix.tobytes("jpeg", quality)
                tw, th = w, h
                orig = len(base["image"])
                if target_w < w and len(jpg) >= orig:
                    # recompress saja tidak untung; turunkan resolusinya (downsample)
                    tw = target_w
                    th = max(1, int(round(h * target_w / w)))
                    mat = fitz.Matrix(tw / w, th / h)
                    small = _resize_pixmap(pix, mat)
                    jpg = small.tobytes("jpeg", quality)
                if len(jpg) < orig:
                    doc.update_stream(xref, jpg, compress=True)
                    for key, val in (("/Filter", "/DCTDecode"), ("/ColorSpace", "/DeviceRGB"), ("/BitsPerComponent", "8")):
                        doc.xref_set_key(xref, key.lstrip("/"), val)
                    doc.xref_set_key(xref, "Width", str(tw))
                    doc.xref_set_key(xref, "Height", str(th))
        doc.save(str(dst), garbage=4, deflate=True, clean=True)
    finally:
        doc.close()
    if dst.stat().st_size >= src.stat().st_size:
        shutil.copyfile(src, dst)


# level -> sisi terpanjang maksimum (px) untuk gambar di dalam DOCX/PPTX.
# Kualitas JPEG memakai angka yang sama dengan COMPRESS_LEVELS.
OOXML_MAX_SIDE = {"light": 2400, "medium": 1600, "max": 1024}
OOXML_MEDIA_DIRS = ("word/media/", "ppt/media/")
OOXML_MAX_PIXELS = 60_000_000  # di atas ini gambar dilewati (anti decompression bomb)


def _shrink_image(data: bytes, ext: str, quality: int, max_side: int) -> bytes | None:
    """Kecilkan satu gambar. Format dan nama tetap, supaya relasi di dokumen tidak putus.

    Return byte baru, atau None kalau gambar dilewati / tidak jadi lebih kecil.
    """
    import io

    from PIL import Image

    Image.MAX_IMAGE_PIXELS = OOXML_MAX_PIXELS
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception:  # rusak, terlalu besar, atau format tidak dikenal: biarkan apa adanya
        return None
    if getattr(img, "is_animated", False):
        return None

    try:
        w, h = img.size
        longest = max(w, h)
        if longest > max_side:
            scale = max_side / longest
            size = (max(1, round(w * scale)), max(1, round(h * scale)))
            if img.mode == "P":  # resize LANCZOS tidak mendukung palette
                img = img.convert("RGBA")
            img = img.resize(size, Image.LANCZOS)

        out = io.BytesIO()
        if ext in ("jpg", "jpeg"):
            if img.mode not in ("RGB", "L", "CMYK"):
                img = img.convert("RGB")
            extra = {}
            # simpan EXIF (termasuk orientasi) dan profil warna supaya tampilan sama persis
            if exif := img.info.get("exif"):
                extra["exif"] = exif
            if icc := img.info.get("icc_profile"):
                extra["icc_profile"] = icc
            img.save(out, "JPEG", quality=quality, optimize=True, **extra)
        else:  # png: tanpa kehilangan data, alpha tetap
            img.save(out, "PNG", optimize=True)
    except Exception:  # mode gambar yang tidak didukung: biarkan apa adanya
        return None

    result = out.getvalue()
    return result if len(result) < len(data) else None


def _compress_ooxml(src: Path, dst: Path, options: dict | None) -> None:
    """Kecilkan DOCX/PPTX dengan memperkecil gambar di word/media atau ppt/media.

    Teks, layout, dan nama file media tidak berubah. Kalau hasil tidak lebih kecil
    dari aslinya (misalnya dokumen tanpa gambar), file asli yang disalin ke output.
    """
    import zipfile

    level = (options or {}).get("level", "medium")
    quality, _ = COMPRESS_LEVELS.get(level, COMPRESS_LEVELS["medium"])
    max_side = OOXML_MAX_SIDE.get(level, OOXML_MAX_SIDE["medium"])
    try:
        zin = zipfile.ZipFile(src)
    except zipfile.BadZipFile:
        raise ConvertError("File rusak dan tidak bisa dibuka. Coba simpan ulang dari Office lalu unggah lagi.")
    with zin, zipfile.ZipFile(dst, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info)
            name = info.filename
            ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
            if name.startswith(OOXML_MEDIA_DIRS) and ext in ("jpg", "jpeg", "png"):
                new = _shrink_image(data, ext, quality, max_side)
                if new is not None:
                    data = new
            # pertahankan nama, tanggal, dan metode kompresi entri aslinya
            clone = zipfile.ZipInfo(name, info.date_time)
            clone.compress_type = info.compress_type
            clone.external_attr = info.external_attr
            zout.writestr(clone, data, compresslevel=9 if info.compress_type == zipfile.ZIP_DEFLATED else None)
    with zipfile.ZipFile(dst) as check:  # hasil harus bisa dibuka dan utuh
        if check.testzip() is not None:
            raise ConvertError("Gagal membuat file hasil. Coba file lain.")
    if dst.stat().st_size >= src.stat().st_size:
        shutil.copyfile(src, dst)


def compress_docx(src: Path, dst: Path, options: dict | None = None) -> None:
    _compress_ooxml(src, dst, options)


def compress_pptx(src: Path, dst: Path, options: dict | None = None) -> None:
    _compress_ooxml(src, dst, options)


def pdf_to_docx(src: Path, dst: Path) -> None:
    if not _has_text(src):
        raise ConvertError(SCAN_MESSAGE)
    from pdf2docx import Converter

    cv = Converter(str(src))
    try:
        cv.convert(str(dst))
    finally:
        cv.close()


def _soffice_convert(src: Path, dst: Path, output_format: str, output_extension: str) -> None:
    exe = os.environ.get("KONVERTIN_SOFFICE_PATH") or shutil.which("soffice") or shutil.which("libreoffice")
    if not exe and os.name == "nt":
        for variable in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
            program_files = os.environ.get(variable)
            candidate = Path(program_files) / "LibreOffice" / "program" / "soffice.exe" if program_files else None
            if candidate and candidate.is_file():
                exe = str(candidate)
                break
    if not exe:
        raise ConvertError(
            "LibreOffice belum terpasang atau tidak ditemukan di server. Konversi format Office ini memerlukan LibreOffice.",
            503,
        )
    with tempfile.TemporaryDirectory() as td:
        td_p = Path(td)
        outdir = td_p / "out"
        outdir.mkdir()
        # salin ke nama aman: hindari karakter aneh & konflik
        tmp_in = td_p / ("in" + src.suffix)
        shutil.copyfile(src, tmp_in)
        cmd = [
            exe, "--headless", "--norestore", "--nolockcheck", "--nodefault", "--nofirststartwizard",
            f"-env:UserInstallation={(td_p / 'profile').as_uri()}",
            "--convert-to", output_format,
            "--outdir", str(outdir), str(tmp_in),
        ]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=110)
        out = outdir / f"in.{output_extension}"
        if r.returncode != 0 or not out.exists():
            raise ConvertError("Dokumen gagal dikonversi. Pastikan file tidak rusak dan format tujuan didukung.")
        shutil.move(str(out), dst)


def _soffice_to_pdf(src: Path, dst: Path) -> None:
    output_format = "pdf:writer_pdf_Export" if src.suffix.lower() == ".docx" else "pdf"
    _soffice_convert(src, dst, output_format, "pdf")


def docx_to_pdf(src: Path, dst: Path) -> None:
    _soffice_to_pdf(src, dst)


def docx_to_odt(src: Path, dst: Path) -> None:
    _soffice_convert(src, dst, "odt", "odt")


def pptx_to_pdf(src: Path, dst: Path) -> None:
    _soffice_to_pdf(src, dst)


def pptx_to_odp(src: Path, dst: Path) -> None:
    _soffice_convert(src, dst, "odp", "odp")


def xlsx_to_pdf(src: Path, dst: Path) -> None:
    _soffice_to_pdf(src, dst)


def xlsx_to_ods(src: Path, dst: Path) -> None:
    _soffice_convert(src, dst, "ods", "ods")


def pdf_to_txt(src: Path, dst: Path) -> None:
    if not _has_text(src):
        raise ConvertError(SCAN_MESSAGE)
    exe = shutil.which("pdftotext")
    if exe:
        r = subprocess.run([exe, "-layout", "-enc", "UTF-8", str(src), str(dst)], capture_output=True, timeout=100)
        if r.returncode == 0 and dst.exists():
            return
    from pypdf import PdfReader

    reader = PdfReader(str(src))
    dst.write_text("\n\n".join((p.extract_text() or "") for p in reader.pages), encoding="utf-8")


def docx_to_txt(src: Path, dst: Path) -> None:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(str(src))
    except Exception:
        raise ConvertError("DOCX rusak dan tidak bisa dibaca. Coba simpan ulang dari Word lalu unggah lagi.")
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(child, doc).text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                lines.append("\t".join(c.text.strip() for c in row.cells))
            lines.append("")
    dst.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def docx_to_md(src: Path, dst: Path) -> None:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(str(src))
    except (OSError, ValueError, zipfile.BadZipFile, PackageNotFoundError):
        raise ConvertError("DOCX rusak dan tidak bisa dibaca. Coba simpan ulang dari Word lalu unggah lagi.")
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            if not text:
                lines.append("")
                continue
            style = paragraph.style.name if paragraph.style else ""
            match = re.fullmatch(r"Heading ([1-6])", style)
            if match:
                lines.append("#" * int(match.group(1)) + " " + text)
            elif "List" in style:
                lines.append("- " + text)
            else:
                lines.append(text)
        elif tag == "tbl":
            rows = ["| " + " | ".join(cell.text.strip().replace("|", "\\|") for cell in row.cells) + " |"
                    for row in Table(child, doc).rows]
            if rows:
                lines.extend((rows[0], "| " + " | ".join("---" for _ in Table(child, doc).rows[0].cells) + " |"))
                lines.extend(rows[1:])
                lines.append("")
    dst.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def txt_to_docx(src: Path, dst: Path) -> None:
    from docx import Document

    try:
        text = src.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        raise ConvertError("File TXT harus menggunakan encoding UTF-8.")
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(str(dst))


_OOXML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_OOXML_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_MAX_SHEET_CELLS = 1_000_000
_MAX_CSV_ROWS = 100_000
_MAX_CSV_COLUMNS = 16_384
_INVALID_XML_CHARS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
ET.register_namespace("", _OOXML_NS)
ET.register_namespace("r", _OOXML_REL_NS)


def _xlsx_column_number(reference: str) -> int:
    match = re.match(r"([A-Z]+)", reference.upper())
    if not match:
        raise ConvertError("Lembar Excel berisi referensi sel yang tidak valid.")
    number = 0
    for character in match.group(1):
        number = number * 26 + ord(character) - ord("A") + 1
    return number


def xlsx_to_csv(src: Path, dst: Path) -> None:
    ns = {"x": _OOXML_NS, "r": _OOXML_REL_NS}
    try:
        with zipfile.ZipFile(src) as package:
            workbook = ET.fromstring(package.read("xl/workbook.xml"))
            relationships = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
            first_sheet = workbook.find("x:sheets/x:sheet", ns)
            if first_sheet is None:
                raise ConvertError("Workbook Excel tidak memiliki lembar kerja.")
            relation_id = first_sheet.get(f"{{{_OOXML_REL_NS}}}id")
            relationship = next(
                (item for item in relationships.findall("x:Relationship", {"x": _PKG_REL_NS})
                 if item.get("Id") == relation_id),
                None,
            )
            if relationship is None or not relationship.get("Target"):
                raise ConvertError("Lembar pertama workbook Excel tidak dapat ditemukan.")
            target = relationship.get("Target", "")
            sheet_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            if not sheet_path.startswith("xl/") or ".." in Path(sheet_path).parts:
                raise ConvertError("Workbook Excel berisi lokasi lembar kerja yang tidak valid.")
            sheet = ET.fromstring(package.read(sheet_path))
            shared_strings: list[str] = []
            if "xl/sharedStrings.xml" in package.namelist():
                shared_root = ET.fromstring(package.read("xl/sharedStrings.xml"))
                shared_strings = [
                    "".join(node.text or "" for node in item.iter(f"{{{_OOXML_NS}}}t"))
                    for item in shared_root.findall("x:si", ns)
                ]
    except ConvertError:
        raise
    except (zipfile.BadZipFile, KeyError, ET.ParseError, ValueError, OSError):
        raise ConvertError("Workbook XLSX rusak atau tidak didukung. Coba simpan ulang di Excel.")

    count = 0
    previous_row = 0
    with dst.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output)
        for row in sheet.findall(".//x:sheetData/x:row", ns):
            try:
                row_number = int(row.get("r", "0"))
            except ValueError:
                raise ConvertError("Lembar Excel memiliki nomor baris yang tidak valid.")
            if row_number < 1 or row_number > 1_048_576 or row_number - previous_row > _MAX_CSV_ROWS:
                raise ConvertError("Lembar Excel melebihi batas baris yang aman untuk CSV.")
            for _ in range(previous_row + 1, row_number):
                writer.writerow([])
            previous_row = row_number
            values: dict[int, str] = {}
            max_column = 0
            for cell in row.findall("x:c", ns):
                column = _xlsx_column_number(cell.get("r", ""))
                if column > _MAX_CSV_COLUMNS:
                    raise ConvertError("Lembar Excel melebihi batas jumlah kolom yang aman untuk CSV.")
                kind = cell.get("t", "")
                value_node = cell.find("x:v", ns)
                value = value_node.text if value_node is not None and value_node.text is not None else ""
                if kind == "s" and value:
                    try:
                        value = shared_strings[int(value)]
                    except (ValueError, IndexError):
                        raise ConvertError("Workbook XLSX memiliki indeks teks yang tidak valid.")
                elif kind == "inlineStr":
                    inline = cell.find("x:is", ns)
                    value = "".join(node.text or "" for node in inline.iter(f"{{{_OOXML_NS}}}t")) if inline is not None else ""
                elif kind == "b":
                    value = "TRUE" if value == "1" else "FALSE"
                values[column] = value
                max_column = max(max_column, column)
                count += 1
                if count > _MAX_SHEET_CELLS:
                    raise ConvertError("Lembar Excel terlalu besar untuk dikonversi dengan aman.")
            writer.writerow([values.get(column, "") for column in range(1, max_column + 1)])
        if previous_row == 0:
            writer.writerow([])


def _column_name(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result


def csv_to_xlsx(src: Path, dst: Path) -> None:
    try:
        text = src.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        raise ConvertError("File CSV harus menggunakan encoding UTF-8.")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|") if sample else csv.excel
    except csv.Error:
        dialect = csv.excel

    ET.register_namespace("", _CONTENT_NS)
    types = ET.Element(f"{{{_CONTENT_NS}}}Types")
    ET.SubElement(types, f"{{{_CONTENT_NS}}}Default", {
        "Extension": "rels", "ContentType": "application/vnd.openxmlformats-package.relationships+xml",
    })
    ET.SubElement(types, f"{{{_CONTENT_NS}}}Default", {
        "Extension": "xml", "ContentType": "application/xml",
    })
    ET.SubElement(types, f"{{{_CONTENT_NS}}}Override", {
        "PartName": "/xl/workbook.xml",
        "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    })
    ET.SubElement(types, f"{{{_CONTENT_NS}}}Override", {
        "PartName": "/xl/worksheets/sheet1.xml",
        "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
    })
    package_rels = ET.Element(f"{{{_PKG_REL_NS}}}Relationships")
    ET.SubElement(package_rels, f"{{{_PKG_REL_NS}}}Relationship", {
        "Id": "rId1",
        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument",
        "Target": "xl/workbook.xml",
    })
    workbook = ET.Element(f"{{{_OOXML_NS}}}workbook")
    sheets = ET.SubElement(workbook, f"{{{_OOXML_NS}}}sheets")
    ET.SubElement(sheets, f"{{{_OOXML_NS}}}sheet", {
        "name": "Sheet1", "sheetId": "1", f"{{{_OOXML_REL_NS}}}id": "rId1",
    })
    workbook_rels = ET.Element(f"{{{_PKG_REL_NS}}}Relationships")
    ET.SubElement(workbook_rels, f"{{{_PKG_REL_NS}}}Relationship", {
        "Id": "rId1",
        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
        "Target": "worksheets/sheet1.xml",
    })
    worksheet = ET.Element(f"{{{_OOXML_NS}}}worksheet")
    sheet_data = ET.SubElement(worksheet, f"{{{_OOXML_NS}}}sheetData")
    xml_space = "{http://www.w3.org/XML/1998/namespace}space"
    cell_count = 0
    try:
        reader = csv.reader(io.StringIO(text, newline=""), dialect)
        for row_number, values in enumerate(reader, start=1):
            if row_number > _MAX_CSV_ROWS or len(values) > _MAX_CSV_COLUMNS:
                raise ConvertError("CSV melebihi batas baris atau kolom yang aman.")
            cell_count += len(values)
            if cell_count > _MAX_SHEET_CELLS:
                raise ConvertError("CSV terlalu besar untuk dikonversi dengan aman.")
            row = ET.SubElement(sheet_data, f"{{{_OOXML_NS}}}row", {"r": str(row_number)})
            for column_number, value in enumerate(values, start=1):
                if len(value) > 32_767 or _INVALID_XML_CHARS.search(value):
                    raise ConvertError("CSV berisi teks yang terlalu panjang atau karakter yang tidak dapat disimpan di XLSX.")
                cell = ET.SubElement(row, f"{{{_OOXML_NS}}}c", {
                    "r": f"{_column_name(column_number)}{row_number}", "t": "inlineStr",
                })
                inline = ET.SubElement(cell, f"{{{_OOXML_NS}}}is")
                text_node = ET.SubElement(inline, f"{{{_OOXML_NS}}}t")
                if value[:1].isspace() or value[-1:].isspace():
                    text_node.set(xml_space, "preserve")
                text_node.text = value
    except csv.Error:
        raise ConvertError("Format CSV tidak valid.")

    try:
        with zipfile.ZipFile(dst, "w", compression=zipfile.ZIP_DEFLATED) as package:
            package.writestr("[Content_Types].xml", ET.tostring(types, encoding="utf-8", xml_declaration=True))
            package.writestr("_rels/.rels", ET.tostring(package_rels, encoding="utf-8", xml_declaration=True))
            package.writestr("xl/workbook.xml", ET.tostring(workbook, encoding="utf-8", xml_declaration=True))
            package.writestr("xl/_rels/workbook.xml.rels", ET.tostring(workbook_rels, encoding="utf-8", xml_declaration=True))
            package.writestr("xl/worksheets/sheet1.xml", ET.tostring(worksheet, encoding="utf-8", xml_declaration=True))
    except OSError:
        raise ConvertError("File XLSX tidak dapat dibuat. Coba CSV yang lebih kecil.")


def pptx_to_txt(src: Path, dst: Path) -> None:
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError

    try:
        presentation = Presentation(str(src))
    except (OSError, ValueError, zipfile.BadZipFile, PackageNotFoundError):
        raise ConvertError("PPTX rusak dan tidak bisa dibaca. Coba simpan ulang di PowerPoint.")
    slides: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines = [f"Slide {number}"]
        for shape in slide.shapes:
            if shape.has_table:
                lines.extend("\t".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows)
            elif shape.has_text_frame and shape.text.strip():
                lines.append(shape.text.strip())
        slides.append("\n".join(lines))
    dst.write_text("\n\n".join(slides) + "\n", encoding="utf-8")


HANDLERS = {
    "pdf_to_docx": pdf_to_docx,
    "docx_to_pdf": docx_to_pdf,
    "docx_to_odt": docx_to_odt,
    "docx_to_md": docx_to_md,
    "pptx_to_pdf": pptx_to_pdf,
    "pptx_to_odp": pptx_to_odp,
    "pptx_to_txt": pptx_to_txt,
    "xlsx_to_pdf": xlsx_to_pdf,
    "xlsx_to_ods": xlsx_to_ods,
    "xlsx_to_csv": xlsx_to_csv,
    "csv_to_xlsx": csv_to_xlsx,
    "pdf_to_txt": pdf_to_txt,
    "docx_to_txt": docx_to_txt,
    "txt_to_docx": txt_to_docx,
    "compress_pdf": compress_pdf,
    "compress_docx": compress_docx,
    "compress_pptx": compress_pptx,
}


def main(argv: list[str]) -> int:
    job_type, src, dst = argv[1], Path(argv[2]), Path(argv[3])
    options: dict = {}
    if len(argv) > 4:
        try:
            options = json.loads(argv[4])
        except (ValueError, TypeError):
            options = {}
    try:
        handler = HANDLERS[job_type]
        if "options" in inspect.signature(handler).parameters:
            handler(src, dst, options)
        else:
            handler(src, dst)
        if not dst.exists() or dst.stat().st_size == 0:
            raise ConvertError("Konversi selesai tapi hasilnya kosong. Coba file lain.")
        return 0
    except ConvertError as e:
        print("USER:" + e.message, file=sys.stderr)
        return 2
    except subprocess.TimeoutExpired:
        print("USER:Konversi terlalu lama dan dihentikan. Coba file yang lebih kecil.", file=sys.stderr)
        return 3
    except MemoryError:
        print("USER:File terlalu berat diproses. Coba file yang lebih kecil.", file=sys.stderr)
        return 3
    except Exception as e:  # jangan bocorkan isi file ke log
        print(f"INTERNAL:{type(e).__name__}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
